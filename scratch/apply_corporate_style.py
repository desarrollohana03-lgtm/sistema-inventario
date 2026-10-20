"""Aplica el estilo corporativo tablero (flat) al reporte XLSX de Estadísticas.

Reglas estrictas:
  - Sin gridlines ni bordes internos entre valores.
  - Una sola línea horizontal gris BAJO la fila de encabezado.
  - Una sola línea horizontal gris SOBRE la fila TOTAL CONSOLIDADO + relleno.
  - Columna "Saldo Fin": valores de datos en rojo corporativo.
  - No se alteran anchos de columna ni alturas de fila.
"""

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side

HANA_RED = "C02026"
BORDER_GRAY = "D9D9D9"
TOTAL_FILL = "EFEFEF"

TOTAL_LABELS = {"TOTAL CONSOLIDADO", "TOTAL CONSULDIDO"}
HEADER_LABEL = "descripción"
SALDO_LABEL = "saldo fin"

GRAY_SIDE = Side(style="thin", color=BORDER_GRAY)
NO_BORDER = Border()


def _norm(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _find_row(ws, predicate, max_col=1):
    for row in ws.iter_rows(min_col=1, max_col=max_col):
        if predicate(_norm(row[0].value)):
            return row[0].row
    return None


def apply_corporate_style(file_path: str, sheet_name: str | None = None) -> str:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo: {path}")

    wb = load_workbook(path)
    ws = wb[sheet_name] if sheet_name else wb.active

    ws.sheet_view.showGridLines = False

    header_row = _find_row(ws, lambda s: s.lower() == HEADER_LABEL)
    if header_row is None:
        header_row = _find_row(
            ws, lambda s: s.lower() in {HEADER_LABEL, "saldo ini"}
        )
    total_row = _find_row(ws, lambda s: s.upper() in TOTAL_LABELS)

    if header_row is None:
        raise ValueError("No se pudo localizar la fila de encabezado.")
    if total_row is None:
        raise ValueError("No se pudo localizar la fila TOTAL CONSOLIDADO.")

    saldo_col = None
    for c in range(1, ws.max_column + 1):
        if _norm(ws.cell(header_row, c).value).lower() == SALDO_LABEL:
            saldo_col = c
    if saldo_col is None and ws.max_column:
        saldo_col = ws.max_column

    for row in ws.iter_rows():
        for cell in row:
            cell.border = NO_BORDER

    for cell in ws[header_row]:
        cell.border = Border(bottom=GRAY_SIDE)

    for cell in ws[total_row]:
        cell.border = Border(top=GRAY_SIDE)
        cell.fill = PatternFill(fill_type="solid", fgColor=TOTAL_FILL)

    if saldo_col is not None:
        for r in range(header_row + 1, total_row):
            cell = ws.cell(r, saldo_col)
            if cell.value is None:
                continue
            cur = cell.font
            cell.font = Font(
                name=cur.name,
                size=cur.size,
                bold=cur.bold,
                italic=cur.italic,
                color=HANA_RED,
            )

    wb.save(path)
    return str(path)


def main() -> int:
    targets = [
        (r"backend\stats_backend.xlsx", None),
    ]
    base = Path(r"C:\Users\Usuario\Desktop\sistema-inventario")
    for rel, sheet in targets:
        full = base / rel
        saved = apply_corporate_style(str(full), sheet)
        sys.stdout.buffer.write((f"OK -> {saved}\n").encode("utf-8"))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
