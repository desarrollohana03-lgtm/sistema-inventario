import sys
import openpyxl
from openpyxl.styles import Border, Side, PatternFill, Font
from openpyxl.utils import get_column_letter

PATH = r"C:\Users\Usuario\Desktop\sistema-inventario\backend\stats_backend.xlsx"

HANA_RED = "C02026"
LINE_GRAY = "A6A6A6"
TOTAL_FILL = "EFEFEF"

HEADER_ROW = 6
SALDO_FIN_HEADER = "Saldo Fin"


def find_total_row(ws):
    for r in range(ws.max_row, 0, -1):
        val = ws.cell(row=r, column=1).value
        if val is None:
            continue
        text = str(val).strip().upper().replace(" ", "")
        if text in ("TOTALCONSOLIDADO", "TOTALCONSULDIDO", "TOTALCONSOLIDADO"):
            return r
        if text.startswith("TOTAL"):
            return r
    return None


def main():
    wb = openpyxl.load_workbook(PATH)
    ws = wb.active

    ws.sheet_view.showGridLines = False

    max_col = ws.max_column
    total_row = find_total_row(ws)
    if total_row is None:
        print("ERROR: could not locate TOTAL row")
        sys.exit(1)

    first_data_row = HEADER_ROW + 1
    last_data_row = total_row - 1

    saldo_fin_col = None
    for c in range(1, max_col + 1):
        if str(ws.cell(row=HEADER_ROW, column=c).value).strip() == SALDO_FIN_HEADER:
            saldo_fin_col = c
            break
    if saldo_fin_col is None:
        saldo_fin_col = max_col

    none_side = Side(style=None)
    no_border = Border(left=none_side, right=none_side, top=none_side, bottom=none_side)

    bottom_gray = Border(bottom=Side(style="thin", color=LINE_GRAY),
                         left=none_side, right=none_side, top=none_side)
    top_gray = Border(top=Side(style="thin", color=LINE_GRAY),
                      left=none_side, right=none_side, bottom=none_side)

    total_fill = PatternFill(fill_type="solid", fgColor=TOTAL_FILL)
    red_font = Font(color=HANA_RED)

    for r in range(first_data_row, last_data_row + 1):
        for c in range(1, max_col + 1):
            ws.cell(row=r, column=c).border = no_border

    for r in (1, 2, 3, 4, 5):
        for c in range(1, max_col + 1):
            ws.cell(row=r, column=c).border = no_border

    for c in range(1, max_col + 1):
        ws.cell(row=HEADER_ROW, column=c).border = bottom_gray

    for c in range(1, max_col + 1):
        cell = ws.cell(row=total_row, column=c)
        cell.border = top_gray
        cell.fill = total_fill

    for r in range(first_data_row, last_data_row + 1):
        cell = ws.cell(row=r, column=saldo_fin_col)
        cell.font = Font(
            name=cell.font.name,
            sz=cell.font.sz,
            b=cell.font.b,
            i=cell.font.i,
            color=HANA_RED,
        )

    wb.save(PATH)

    print("OK")
    print("file:", PATH)
    print("sheet:", ws.title)
    print("max_col:", max_col, "| header_row:", HEADER_ROW,
          "| data:", first_data_row, "-", last_data_row,
          "| total_row:", total_row, "| saldo_fin_col:", get_column_letter(saldo_fin_col))
    print("showGridLines:", ws.sheet_view.showGridLines)


if __name__ == "__main__":
    main()
